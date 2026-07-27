"""
1단계: 텍스트 추출기

원본 파일에서 텍스트를 추출하여 LangChain Document 객체 리스트로 반환합니다.
- notice_AICA.docx → 공지사항 단위로 분리 → Document 1개씩 생성
- 첨부파일 (PDF / DOCX / PNG / JPG) → Document 생성
- 규정 문서 (PDF / MD) → Document 생성

실행 방법:
    python notice_extractor.py

사전 설치:
    pip install langchain python-docx pymupdf pytesseract pillow
    (Windows) Tesseract OCR 설치 필요
    https://github.com/UB-Mannheim/tesseract/wiki
"""

import os
import re
import logging

import fitz                              # PDF 추출 (pymupdf)
from docx import Document as DocxFile
from PIL import Image

import pytesseract
pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

from langchain_core.documents import Document

import config


# ───────────────────────────────────────────
# 로깅 설정
# ───────────────────────────────────────────
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.StreamHandler(),
        logging.FileHandler("extract_log.txt", encoding="utf-8")
    ]
)
logger = logging.getLogger(__name__)


# ───────────────────────────────────────────
# 공통 유틸
# ───────────────────────────────────────────
def normalize_date(raw: str) -> str:
    """
    날짜 문자열을 YYYY-MM-DD로 변환
    예: 26/05/08 → 2026-05-08
    """
    parts = re.split(r"[/\-]", raw.strip())
    if len(parts) != 3:
        return raw
    year, month, day = parts
    if len(year) == 2:
        year = "20" + year
    return f"{year}-{month.zfill(2)}-{day.zfill(2)}"


def is_notice_title(para) -> bool:
    """Word 자동 번호 목록(ilvl=0) 단락인지 확인 → 공지 제목 판별"""
    NS    = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    ilvl  = para._element.find(f".//{{{NS}}}ilvl")
    numId = para._element.find(f".//{{{NS}}}numId")
    if ilvl is None or numId is None:
        return False
    return ilvl.get(f"{{{NS}}}val") == "0"


def parse_title_and_date(text: str):
    """
    제목 단락에서 제목과 날짜를 분리
    예: "출석인정 안내 2026/05/06" → ("출석인정 안내", "2026-05-06")
    """
    match = re.search(r"(\d{2,4}[/\-]\d{2}[/\-]\d{2})\s*$", text)
    if match:
        title = text[:match.start()].strip()
        return title, normalize_date(match.group(1))
    return text.strip(), None


# ───────────────────────────────────────────
# 1. 공지사항 단위 분리 (notice_AICA.docx)
# ───────────────────────────────────────────
def extract_notices_from_docx(docx_path: str) -> list[Document]:
    """
    notice_AICA.docx를 공지사항 단위로 분리하여 Document 리스트 반환
    공지사항 1개 = Document 1개
    """
    doc = DocxFile(docx_path)
    documents = []
    current = None
    notice_index = 0  # 문서 내 순서 (1부터 시작)

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        if is_notice_title(para):
            # 이전 공지 저장
            if current:
                documents.append(_build_notice_document(current))

            notice_index += 1
            title, release_date = parse_title_and_date(text)
            current = {
                "notice_index" : notice_index,
                "title"        : title,
                "release_date" : release_date,
                "content"      : "",
                "attachments"  : []
            }

        elif current:
            attachment_match = re.match(r"^\[첨부\]\s+(.+)$", text)
            if attachment_match:
                current["attachments"].append(attachment_match.group(1).strip())
            else:
                current["content"] += text + "\n"

    # 마지막 공지 저장
    if current:
        documents.append(_build_notice_document(current))

    logger.info(f"[공지] 분리 완료: {len(documents)}건")
    return documents


def _build_notice_document(notice: dict) -> Document:
    """공지사항 딕셔너리를 Document 객체로 변환"""
    content = notice["content"].strip()
    attachments = notice["attachments"]

    # page_content: 제목 + 날짜 + 본문
    page_content = f"제목: {notice['title']}\n날짜: {notice['release_date'] or '날짜 없음'}\n\n{content}"

    # 첨부파일 목록이 있으면 본문에 추가
    if attachments:
        page_content += "\n\n[첨부파일]\n" + "\n".join(attachments)

    return Document(
        page_content=page_content,
        metadata={
            "source_type"  : "notice",
            "notice_index" : notice["notice_index"],
            "title"        : notice["title"],
            "release_date" : notice["release_date"] or "",
            "attachments"  : ", ".join(attachments)  # Pinecone 메타데이터는 문자열로
        }
    )


# ───────────────────────────────────────────
# 2. 파일 형식별 텍스트 추출
# ───────────────────────────────────────────
def extract_text_from_pdf(file_path: str) -> str:
    """PDF에서 텍스트 추출"""
    text = ""
    with fitz.open(file_path) as pdf:
        for page in pdf:
            text += page.get_text()
    return text.strip()


def extract_text_from_docx(file_path: str) -> str:
    """DOCX에서 텍스트 추출"""
    doc = DocxFile(file_path)
    return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])


def extract_text_from_image(file_path: str) -> str:
    """PNG/JPG에서 OCR로 텍스트 추출 (한국어 + 영어)"""
    img = Image.open(file_path)
    return pytesseract.image_to_string(img, lang="kor+eng").strip()


def extract_text_from_file(file_path: str) -> str:
    """확장자에 따라 적절한 추출 함수 호출"""
    ext = os.path.splitext(file_path)[-1].lower()

    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    elif ext in (".png", ".jpg", ".jpeg"):
        return extract_text_from_image(file_path)
    elif ext == ".md":
        with open(file_path, encoding="utf-8") as f:
            return f.read().strip()
    else:
        logger.warning(f"[SKIP] 지원하지 않는 형식: {file_path}")
        return ""


# ───────────────────────────────────────────
# 3. 공지사항 첨부파일 추출
# ───────────────────────────────────────────
def extract_from_attachments(attachment_dir: str) -> list[Document]:
    """
    notice_attachment 폴더를 순회하여 첨부파일에서 Document 생성
    폴더명 = notice_id
    """
    documents = []

    for folder_name in sorted(os.listdir(attachment_dir)):
        folder_path = os.path.join(attachment_dir, folder_name)
        if not os.path.isdir(folder_path):
            continue

        for file_name in os.listdir(folder_path):
            file_path = os.path.join(folder_path, file_name)
            if not os.path.isfile(file_path):
                continue

            text = extract_text_from_file(file_path)
            if not text:
                logger.warning(f"[SKIP] 텍스트 없음: {file_name}")
                continue

            file_type = os.path.splitext(file_name)[-1].lstrip(".").lower()
            doc = Document(
                page_content=text,
                metadata={
                    "source_type" : "attachment",
                    "notice_id"   : folder_name,   # 폴더명 = notice_id
                    "file_name"   : file_name,
                    "file_type"   : file_type,
                }
            )
            documents.append(doc)
            logger.info(f"[첨부] 추출 완료: {file_name} (notice_id={folder_name})")

    logger.info(f"[첨부] 전체 추출 완료: {len(documents)}건")
    return documents


# ───────────────────────────────────────────
# 4. 규정 문서 추출
# ───────────────────────────────────────────
def extract_from_rules(rule_dir: str) -> list[Document]:
    """rule 폴더의 문서에서 Document 생성"""
    documents = []

    for file_name in os.listdir(rule_dir):
        file_path = os.path.join(rule_dir, file_name)
        if not os.path.isfile(file_path):
            continue

        text = extract_text_from_file(file_path)
        if not text:
            logger.warning(f"[SKIP] 텍스트 없음: {file_name}")
            continue

        file_type = os.path.splitext(file_name)[-1].lstrip(".").lower()
        doc = Document(
            page_content=text,
            metadata={
                "source_type" : "reference",
                "file_name"   : file_name,
                "file_type"   : file_type,
            }
        )
        documents.append(doc)
        logger.info(f"[규정] 추출 완료: {file_name}")

    logger.info(f"[규정] 전체 추출 완료: {len(documents)}건")
    return documents


# ───────────────────────────────────────────
# 5. 전체 실행
# ───────────────────────────────────────────
def extract_all() -> list[Document]:
    """전체 데이터 소스에서 Document 추출 후 통합 반환"""
    logger.info("=== 텍스트 추출 시작 ===")

    notice_docs     = extract_notices_from_docx(config.NOTICE_DOCX)
    attachment_docs = extract_from_attachments(config.ATTACHMENT_DIR)
    rule_docs       = extract_from_rules(config.RULE_DIR)

    all_docs = notice_docs + attachment_docs + rule_docs

    logger.info(
        f"=== 추출 완료 | "
        f"공지:{len(notice_docs)} "
        f"첨부:{len(attachment_docs)} "
        f"규정:{len(rule_docs)} "
        f"합계:{len(all_docs)} ==="
    )
    return all_docs


if __name__ == "__main__":
    docs = extract_all()

    print(f"\n총 Document 수: {len(docs)}개")
    print("\n--- 첫 번째 Document 샘플 ---")
    print(f"[Metadata]\n{docs[0].metadata}")
    print(f"\n[page_content 미리보기]\n{docs[0].page_content[:300]}")
    